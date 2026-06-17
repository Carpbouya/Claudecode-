import csv
import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import DATA_DIR
from storage import get_connection


LISTING_COLUMNS = [
    ("求人タイトル", "title"),
    ("企業名", "company_name"),
    ("年収下限", "salary_min"),
    ("年収上限", "salary_max"),
    ("年収テキスト", "salary_text"),
    ("勤務地", "location"),
    ("都道府県", "prefecture"),
    ("職種", "job_category"),
    ("概要", "description_summary"),
    ("求人URL", "source_url"),
    ("取得日時", "scraped_at"),
]

DETAIL_COLUMNS = LISTING_COLUMNS[:8] + [
    ("必須スキル・資格", "required_skills"),
    ("歓迎スキル", "preferred_skills"),
    ("経験年数", "experience_years"),
    ("雇用形態", "employment_type"),
    ("業界", "industry"),
    ("勤務形態", "work_style"),
    ("福利厚生", "benefits"),
    ("求人URL", "source_url"),
    ("取得日時", "scraped_at"),
]


def export_csv(filename: str | None = None, with_details: bool = False) -> str:
    os.makedirs(DATA_DIR, exist_ok=True)
    if not filename:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        suffix = "_details" if with_details else ""
        filename = f"jobs{suffix}_{ts}.csv"

    path = os.path.join(DATA_DIR, filename)
    conn = get_connection()

    if with_details:
        columns = DETAIL_COLUMNS
        rows = conn.execute(
            """SELECT j.*, d.required_skills, d.preferred_skills,
                      d.experience_years, d.employment_type, d.industry,
                      d.work_style, d.benefits
               FROM jobs j JOIN job_details d ON j.id = d.job_id
               ORDER BY j.prefecture, j.salary_max DESC"""
        ).fetchall()
    else:
        columns = LISTING_COLUMNS
        rows = conn.execute(
            "SELECT * FROM jobs ORDER BY prefecture, salary_max DESC"
        ).fetchall()
    conn.close()

    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow([c[0] for c in columns])
        for row in rows:
            row_dict = dict(row)
            writer.writerow([
                _format_value(row_dict.get(c[1], ""), c[1]) for c in columns
            ])

    return path


def export_excel(filename: str | None = None, with_details: bool = False) -> str:
    os.makedirs(DATA_DIR, exist_ok=True)
    if not filename:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        suffix = "_details" if with_details else ""
        filename = f"jobs{suffix}_{ts}.xlsx"

    path = os.path.join(DATA_DIR, filename)
    conn = get_connection()

    if with_details:
        columns = DETAIL_COLUMNS
        rows = conn.execute(
            """SELECT j.*, d.required_skills, d.preferred_skills,
                      d.experience_years, d.employment_type, d.industry,
                      d.work_style, d.benefits
               FROM jobs j JOIN job_details d ON j.id = d.job_id
               ORDER BY j.prefecture, j.salary_max DESC"""
        ).fetchall()
    else:
        columns = LISTING_COLUMNS
        rows = conn.execute(
            "SELECT * FROM jobs ORDER BY prefecture, salary_max DESC"
        ).fetchall()
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "求人データ"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1a237e", end_color="1a237e", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for col_idx, (label, _) in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    for row_idx, row in enumerate(rows, 2):
        row_dict = dict(row)
        for col_idx, (_, key) in enumerate(columns, 1):
            val = _format_value(row_dict.get(key, ""), key)
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border

    for col_idx, (label, _) in enumerate(columns, 1):
        col_letter = get_column_letter(col_idx)
        if "URL" in label:
            ws.column_dimensions[col_letter].width = 40
        elif "タイトル" in label or "スキル" in label:
            ws.column_dimensions[col_letter].width = 35
        elif "企業" in label:
            ws.column_dimensions[col_letter].width = 25
        else:
            ws.column_dimensions[col_letter].width = 15

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"

    wb.save(path)
    return path


def _format_value(val, key: str):
    if val is None:
        return ""
    if key in ("salary_min", "salary_max") and isinstance(val, (int, float)):
        return int(val)
    return str(val)
