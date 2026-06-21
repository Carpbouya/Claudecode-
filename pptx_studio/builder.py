"""
PPTX builder — slide factories + chart helpers.
Each slide type reads from a plain dict and a Theme, writes to a Presentation.
"""
import io
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.font_manager as fm
import numpy as np

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from ds import Theme, hex_rgb

W = 13.33   # slide width  (inches, 16:9)
H = 7.5     # slide height (inches, 16:9)

_JP_FONT = '/usr/share/fonts/opentype/ipafont-gothic/ipag.ttf'

# ── Low-level primitives ───────────────────────────────────────

def _c(h): return hex_rgb(h)

def _rect(slide, l, t, w, h, fill=None, line=None):
    sh = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = _c(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = _c(line)
    else:
        sh.line.fill.background()
    return sh

def _box(slide, l, t, w, h, text, sz,
         bold=False, italic=False, color="#111827", align=PP_ALIGN.LEFT):
    txb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = txb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(sz)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = _c(color)
    return txb

def _multibox(slide, l, t, w, h, lines, sz, color="#111827",
              bullet="▶", line_h=0.5):
    """Render multiple lines, each prefixed with bullet char."""
    txb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = txb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        if i > 0:
            from pptx.util import Pt as _Pt
            p.space_before = _Pt(6)
        run = p.add_run()
        run.text = f"{bullet}  {line}" if bullet else line
        run.font.size = Pt(sz)
        run.font.color.rgb = _c(color)

def _bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _c(color)

def _img(slide, buf, l, t, w, h):
    slide.shapes.add_picture(buf, Inches(l), Inches(t), Inches(w), Inches(h))

def _logo_img(slide, logo_bytes: bytes, l: float, t: float,
              max_h: float, max_w: float | None = None):
    """Add logo preserving aspect ratio; capped to max_h (and max_w if given)."""
    if not logo_bytes:
        return
    try:
        buf = io.BytesIO(logo_bytes)
        sp = slide.shapes.add_picture(buf, Inches(l), Inches(t), height=Inches(max_h))
        if max_w:
            actual_w_in = sp.width / 914400  # EMU → inches
            if actual_w_in > max_w:
                sp._element.getparent().remove(sp._element)
                buf.seek(0)
                slide.shapes.add_picture(buf, Inches(l), Inches(t), width=Inches(max_w))
    except Exception:
        pass

def _title_bar(slide, t: Theme, title: str, subtitle: str = "",
               logo_bytes: bytes | None = None):
    _rect(slide, 0, 0, W, 0.07, fill=t.colors.accent)
    _rect(slide, 0, 0.07, W, 1.2, fill=t.colors.primary)
    title_w = W - 1.1
    if logo_bytes:
        logo_max_h = 0.82
        logo_max_w = 2.2
        logo_l = W - logo_max_w - 0.2
        _logo_img(slide, logo_bytes, logo_l, 0.22, logo_max_h, logo_max_w)
        title_w = logo_l - 0.55 - 0.1
    _box(slide, 0.55, 0.18, title_w, 0.72, title,
         t.typo.sz_h1, bold=True, color=t.colors.text_light)
    if subtitle:
        _box(slide, 0.55, 0.9, title_w, 0.35, subtitle,
             t.typo.sz_sm, color="#9CA3AF")

def _footer(slide, t: Theme, pg: int, company: str = ""):
    _rect(slide, 0, H-0.24, W, 0.24, fill=t.colors.primary)
    if company:
        _box(slide, 0.4, H-0.23, 9.0, 0.22, company,
             8, color="#9CA3AF")
    _box(slide, W-1.3, H-0.23, 1.1, 0.22, str(pg),
         8, color="#9CA3AF", align=PP_ALIGN.RIGHT)

def _blank(prs): return prs.slides.add_slide(prs.slide_layouts[6])

# ── Chart helpers ──────────────────────────────────────────────

def _jp_font():
    try:
        prop = fm.FontProperties(fname=_JP_FONT)
        plt.rcParams['font.family'] = prop.get_name()
    except Exception:
        pass
    plt.rcParams['axes.unicode_minus'] = False

def _parse_csv(txt: str):
    labels, vals = [], []
    for line in txt.strip().splitlines():
        parts = [p.strip() for p in line.split(',')]
        if len(parts) >= 2:
            try:
                labels.append(parts[0])
                vals.append(float(parts[1].replace(',', '').replace('%', '')))
            except ValueError:
                pass
    return labels, vals

def _chart_bar(csv: str, title: str, ylabel: str, accent: str) -> io.BytesIO | None:
    _jp_font()
    labels, vals = _parse_csv(csv)
    if not labels:
        return None
    fig, ax = plt.subplots(figsize=(9.5, 4.2), facecolor='white')
    x = np.arange(len(labels))
    bars = ax.bar(x, vals, color=accent, alpha=0.88, width=0.55, zorder=3, linewidth=0)
    ax.set_facecolor('white')
    ax.set_xticks(x)
    ax.set_xticklabels(labels, fontsize=11)
    ax.yaxis.set_tick_params(labelsize=10)
    for sp in ('top', 'right'):
        ax.spines[sp].set_visible(False)
    ax.spines['left'].set_color('#E5E7EB')
    ax.spines['bottom'].set_color('#E5E7EB')
    ax.yaxis.grid(True, color='#F3F4F6', linewidth=0.8, zorder=0)
    ax.set_axisbelow(True)
    if ylabel:
        ax.set_ylabel(ylabel, fontsize=10, color='#6B7280')
    if title:
        ax.set_title(title, fontsize=12, fontweight='bold', color='#1F2937', pad=10)
    for bar, v in zip(bars, vals):
        ax.text(bar.get_x()+bar.get_width()/2, bar.get_height()+max(vals)*0.012,
                f'{v:,.0f}', ha='center', va='bottom', fontsize=9, color='#374151')
    plt.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format='png', dpi=150, bbox_inches='tight', facecolor='white')
    plt.close(fig)
    buf.seek(0)
    return buf

def _chart_line(csv: str, title: str, ylabel: str, accent: str) -> io.BytesIO | None:
    _jp_font()
    labels, vals = _parse_csv(csv)
    if not labels:
        return None
    fig, ax = plt.subplots(figsize=(9.5, 4.2), facecolor='white')
    x = np.arange(len(labels))
    ax.plot(x, vals, color=accent, linewidth=2.5, zorder=3)
    ax.fill_between(x, vals, alpha=0.10, color=accent)
    ax.scatter(x, vals, color=accent, s=55, zorder=4, linewidth=0)
    ax.set_facecolor('white')
    ax.set_xticks(x)
    ax.set_xticklabels(labels, fontsize=11)
    ax.yaxis.set_tick_params(labelsize=10)
    for sp in ('top', 'right'):
        ax.spines[sp].set_visible(False)
    ax.spines['left'].set_color('#E5E7EB')
    ax.spines['bottom'].set_color('#E5E7EB')
    ax.yaxis.grid(True, color='#F3F4F6', linewidth=0.8, zorder=0)
    ax.set_axisbelow(True)
    if ylabel:
        ax.set_ylabel(ylabel, fontsize=10, color='#6B7280')
    if title:
        ax.set_title(title, fontsize=12, fontweight='bold', color='#1F2937', pad=10)
    plt.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format='png', dpi=150, bbox_inches='tight', facecolor='white')
    plt.close(fig)
    buf.seek(0)
    return buf

def _chart_pie(csv: str, title: str, accent: str) -> io.BytesIO | None:
    _jp_font()
    labels, vals = _parse_csv(csv)
    if not labels:
        return None
    grays = ['#6B7280', '#9CA3AF', '#D1D5DB', '#E5E7EB', '#F3F4F6']
    colors = [accent] + (grays * 4)[:len(labels)-1]
    fig, ax = plt.subplots(figsize=(8, 4.2), facecolor='white')
    wedges, texts, autos = ax.pie(
        vals, labels=labels, colors=colors,
        autopct='%1.1f%%', startangle=90,
        pctdistance=0.75, labeldistance=1.15,
        wedgeprops={'linewidth': 2, 'edgecolor': 'white'},
    )
    for at in autos:
        at.set_fontsize(9)
    for tx in texts:
        tx.set_fontsize(10)
    if title:
        ax.set_title(title, fontsize=12, fontweight='bold', color='#1F2937', pad=10)
    plt.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format='png', dpi=150, bbox_inches='tight', facecolor='white')
    plt.close(fig)
    buf.seek(0)
    return buf

# ── Slide builders ─────────────────────────────────────────────

def s_title(prs, d, t: Theme, pg, company, logo_bytes=None):
    slide = _blank(prs)
    _bg(slide, t.colors.primary)
    # Vertical accent stripe + echo in secondary color
    _rect(slide, 0, 0, 0.09, H, fill=t.colors.accent)
    _rect(slide, 0.09, 0, 0.035, H, fill=t.colors.secondary)
    # Right decorative panel
    _rect(slide, 10.8, 0, 2.53, H, fill=t.colors.secondary)
    # Divider
    _rect(slide, 0.85, 3.88, 5.0, 0.05, fill=t.colors.accent)
    # Logo — prominent placement in upper-right
    if logo_bytes:
        _logo_img(slide, logo_bytes, W-4.2, 0.3, 1.55, 3.5)
    # Title text (shrink if logo present)
    title_w = 9.5 if logo_bytes else 12.0
    _box(slide, 0.85, 1.6, title_w, 2.1,
         d.get("title", "タイトルを入力"),
         t.typo.sz_huge, bold=True, color=t.colors.text_light)
    if d.get("subtitle"):
        _box(slide, 0.85, 3.95, title_w, 0.7, d["subtitle"],
             t.typo.sz_h3, color="#9CA3AF")
    metas = [x for x in [d.get("date"), d.get("author")] if x]
    if metas:
        _box(slide, 0.85, 4.85, 11.0, 0.45, "  ·  ".join(metas),
             t.typo.sz_sm, color="#6B7280")
    if company:
        _box(slide, 9.5, H-0.65, 3.5, 0.5, company,
             t.typo.sz_sm, color="#6B7280", align=PP_ALIGN.RIGHT)


def s_section(prs, d, t: Theme, pg, company, logo_bytes=None):
    slide = _blank(prs)
    _bg(slide, t.colors.surface_alt)
    _rect(slide, 0, 0, 0.5, H, fill=t.colors.accent)
    _box(slide, 1.0, 1.4, 11.0, 0.55,
         f"SECTION  {d.get('section_number', '01')}",
         t.typo.sz_sm, bold=True, color=t.colors.accent)
    _box(slide, 1.0, 1.95, 11.0, 2.0,
         d.get("section_title", ""),
         int(t.typo.sz_h1 * 1.15), bold=True, color=t.colors.text_dark)
    _rect(slide, 1.0, 3.85, 9.5, 0.04, fill=t.colors.divider)
    if d.get("description"):
        _box(slide, 1.0, 4.05, 11.0, 2.8,
             d["description"], t.typo.sz_body, color=t.colors.text_mid)
    _footer(slide, t, pg, company)


def s_kpi3(prs, d, t: Theme, pg, company, logo_bytes=None):
    slide = _blank(prs)
    _bg(slide, t.colors.surface)
    _title_bar(slide, t, d.get("slide_title", ""), d.get("subtitle", ""), logo_bytes)
    card_w, card_h, top = 3.7, 4.9, 1.38
    xs = [0.38, 4.48, 8.58]
    for i in range(1, 4):
        x = xs[i-1]
        _rect(slide, x, top, card_w, card_h, fill=t.colors.surface_alt)
        _rect(slide, x, top, card_w, 0.07, fill=t.colors.accent)
        _box(slide, x+0.22, top+0.22, card_w-0.44, 0.55,
             d.get(f"kpi{i}_label", f"指標{i}"),
             t.typo.sz_body, color=t.colors.text_mid)
        _box(slide, x+0.22, top+0.85, card_w-0.44, 1.6,
             d.get(f"kpi{i}_value", "—"),
             t.typo.sz_kpi, bold=True, color=t.colors.text_dark)
        if d.get(f"kpi{i}_unit"):
            _box(slide, x+0.22, top+2.55, card_w-0.44, 0.42,
                 d[f"kpi{i}_unit"], t.typo.sz_sm, color=t.colors.text_mid)
        if d.get(f"kpi{i}_note"):
            _box(slide, x+0.22, top+3.05, card_w-0.44, 1.6,
                 d[f"kpi{i}_note"], t.typo.sz_sm, color=t.colors.text_mid)
    _footer(slide, t, pg, company)


def s_bullet(prs, d, t: Theme, pg, company, logo_bytes=None):
    slide = _blank(prs)
    _bg(slide, t.colors.surface)
    _title_bar(slide, t, d.get("slide_title", ""), d.get("subtitle", ""), logo_bytes)
    cy = 1.48
    if d.get("lead"):
        _box(slide, 0.55, cy, W-1.1, 0.52,
             d["lead"], t.typo.sz_body+1, italic=True, color=t.colors.text_mid)
        cy += 0.60
    raw = d.get("bullets", "")
    lines = [ln.strip() for ln in raw.splitlines() if ln.strip()]
    for ln in lines:
        _rect(slide, 0.55, cy+0.19, 0.09, 0.09, fill=t.colors.accent)
        _box(slide, 0.82, cy, W-1.45, 0.52,
             ln, t.typo.sz_body, color=t.colors.text_dark)
        cy += 0.54
        if cy > H - 0.55:
            break
    _footer(slide, t, pg, company)


def s_two_column(prs, d, t: Theme, pg, company, logo_bytes=None):
    slide = _blank(prs)
    _bg(slide, t.colors.surface)
    _title_bar(slide, t, d.get("slide_title", ""), d.get("subtitle", ""), logo_bytes)
    mid = 6.78
    col_w = 5.9
    top = 1.42
    ch = H - top - 0.32
    _rect(slide, mid-0.02, top, 0.04, ch, fill=t.colors.divider)
    for side in ("left", "right"):
        x = 0.38 if side == "left" else mid + 0.38
        if d.get(f"{side}_title"):
            _rect(slide, x, top, 0.06, 0.58, fill=t.colors.accent)
            _box(slide, x+0.2, top+0.04, col_w-0.36, 0.56,
                 d[f"{side}_title"], t.typo.sz_h3, bold=True, color=t.colors.text_dark)
        raw = d.get(f"{side}_content", "")
        lines = [ln.strip() for ln in raw.splitlines() if ln.strip()]
        cy = top + 0.74
        for ln in lines:
            _box(slide, x+0.2, cy, col_w-0.36, 0.50,
                 f"•  {ln}", t.typo.sz_body, color=t.colors.text_dark)
            cy += 0.50
            if cy > H - 0.42:
                break
    _footer(slide, t, pg, company)


def _chart_slide(prs, d, t: Theme, pg, company, chart_fn, logo_bytes=None):
    slide = _blank(prs)
    _bg(slide, t.colors.surface)
    _title_bar(slide, t, d.get("slide_title", ""), d.get("subtitle", ""), logo_bytes)
    buf = chart_fn(
        d.get("data_csv", ""), d.get("chart_title", ""),
        d.get("y_label", ""), t.colors.accent,
    )
    if buf:
        _img(slide, buf, 0.45, 1.38, W-0.9, H-2.05)
    if d.get("note"):
        _box(slide, 0.45, H-0.52, W-0.9, 0.3,
             f"※ {d['note']}", t.typo.sz_sm, color=t.colors.text_mid)
    _footer(slide, t, pg, company)


def s_chart_bar(prs, d, t, pg, company, logo_bytes=None):
    _chart_slide(prs, d, t, pg, company, _chart_bar, logo_bytes)


def s_chart_line(prs, d, t, pg, company, logo_bytes=None):
    _chart_slide(prs, d, t, pg, company, _chart_line, logo_bytes)


def s_chart_pie(prs, d, t: Theme, pg, company, logo_bytes=None):
    slide = _blank(prs)
    _bg(slide, t.colors.surface)
    _title_bar(slide, t, d.get("slide_title", ""), d.get("subtitle", ""), logo_bytes)
    buf = _chart_pie(d.get("data_csv", ""), d.get("chart_title", ""), t.colors.accent)
    if buf:
        _img(slide, buf, 1.5, 1.38, W-3.0, H-2.05)
    if d.get("note"):
        _box(slide, 0.45, H-0.52, W-0.9, 0.3,
             f"※ {d['note']}", t.typo.sz_sm, color=t.colors.text_mid)
    _footer(slide, t, pg, company)


def s_table(prs, d, t: Theme, pg, company, logo_bytes=None):
    slide = _blank(prs)
    _bg(slide, t.colors.surface)
    _title_bar(slide, t, d.get("slide_title", ""), d.get("subtitle", ""), logo_bytes)
    raw = d.get("data_csv", "")
    rows = [[c.strip() for c in r.split(',')] for r in raw.strip().splitlines() if r.strip()]
    if not rows:
        _footer(slide, t, pg, company)
        return
    ncols = max(len(r) for r in rows)
    rows = [r + [''] * (ncols - len(r)) for r in rows]
    nrows = len(rows)
    tl, tr = 0.38, 1.42
    tw = W - 0.76
    max_h = H - tr - 0.32
    rh = min(max_h / nrows, 0.58)
    cw = tw / ncols
    sz = max(t.typo.sz_body - 1, 9)
    for ri, row in enumerate(rows):
        for ci, cell in enumerate(row):
            xl = tl + ci * cw
            yt = tr + ri * rh
            if ri == 0:
                fill = t.colors.primary
                tc = t.colors.text_light
            elif ri % 2 == 0:
                fill = t.colors.surface_alt
                tc = t.colors.text_dark
            else:
                fill = t.colors.surface
                tc = t.colors.text_dark
            _rect(slide, xl, yt, cw-0.01, rh-0.01, fill=fill)
            _box(slide, xl+0.1, yt+0.04, cw-0.2, rh-0.06,
                 cell, sz, bold=(ri == 0), color=tc)
    if d.get("note"):
        ny = tr + nrows * rh + 0.08
        if ny < H - 0.35:
            _box(slide, tl, ny, tw, 0.3,
                 f"※ {d['note']}", t.typo.sz_sm, color=t.colors.text_mid)
    _footer(slide, t, pg, company)


def s_quote(prs, d, t: Theme, pg, company, logo_bytes=None):
    slide = _blank(prs)
    _bg(slide, t.colors.accent)
    _box(slide, 0.45, 0.15, 3.0, 2.8, "❝",
         110, color="#FFFFFF18", align=PP_ALIGN.LEFT)
    _box(slide, 1.1, 1.5, W-2.2, 3.6,
         d.get("message", ""),
         t.typo.sz_h1, bold=True, color=t.colors.text_light, align=PP_ALIGN.CENTER)
    if d.get("source"):
        _rect(slide, W/2-2.2, 5.2, 4.4, 0.045, fill="#FFFFFF55")
        _box(slide, 0.5, 5.4, W-1.0, 0.55, d["source"],
             t.typo.sz_sm, color="#FFFFFFCC", align=PP_ALIGN.CENTER)
    _footer(slide, t, pg, company)


def s_closing(prs, d, t: Theme, pg, company, logo_bytes=None):
    slide = _blank(prs)
    _bg(slide, t.colors.primary)
    _rect(slide, 0, 0, W, 0.07, fill=t.colors.accent)
    _box(slide, 0.8, 1.7, W-1.6, 2.1,
         d.get("main_message", "ありがとうございました"),
         t.typo.sz_huge, bold=True, color=t.colors.text_light, align=PP_ALIGN.CENTER)
    _rect(slide, W/2-3.2, 3.8, 6.4, 0.05, fill=t.colors.accent)
    if d.get("contact"):
        _box(slide, 1.5, 4.05, W-3.0, 1.6,
             d["contact"], t.typo.sz_body, color="#9CA3AF", align=PP_ALIGN.CENTER)
    if d.get("next_steps"):
        lines = [ln.strip() for ln in d["next_steps"].splitlines() if ln.strip()]
        cy = 5.25
        for ln in lines:
            _box(slide, 2.0, cy, W-4.0, 0.46,
                 f"▶  {ln}", t.typo.sz_body, color=t.colors.text_light, align=PP_ALIGN.CENTER)
            cy += 0.46
    if company:
        _box(slide, 0.5, H-0.62, W-1.0, 0.42,
             company, t.typo.sz_sm, color="#6B7280", align=PP_ALIGN.CENTER)


# ── Slide metadata (type definitions) ─────────────────────────

SLIDE_TYPES = {
    "title":       {"label": "タイトルスライド",     "icon": "🏠", "desc": "プレゼンの表紙"},
    "section":     {"label": "セクション区切り",     "icon": "📑", "desc": "章の区切り"},
    "kpi3":        {"label": "KPIカード（3つ）",    "icon": "📊", "desc": "大きな数値・指標を3つ並べる"},
    "bullet":      {"label": "箇条書き",            "icon": "📝", "desc": "タイトル＋箇条書きリスト"},
    "two_column":  {"label": "2カラム",             "icon": "⬜", "desc": "左右2列のテキストレイアウト"},
    "chart_bar":   {"label": "棒グラフ",            "icon": "📊", "desc": "棒グラフ（データはラベル,値 形式）"},
    "chart_line":  {"label": "折れ線グラフ",        "icon": "📈", "desc": "トレンドを折れ線で表示"},
    "chart_pie":   {"label": "円グラフ",            "icon": "🥧", "desc": "構成比を円グラフで表示"},
    "table":       {"label": "データテーブル",      "icon": "📋", "desc": "CSVからテーブルを自動生成"},
    "quote":       {"label": "キーメッセージ",      "icon": "💬", "desc": "インパクトのある一言を大きく"},
    "closing":     {"label": "クロージング",        "icon": "🎯", "desc": "まとめ・お問い合わせ"},
}

SLIDE_FIELDS = {
    "title": [
        {"key": "title",    "label": "メインタイトル",   "type": "text",     "default": "タイトルを入力"},
        {"key": "subtitle", "label": "サブタイトル",     "type": "text",     "default": ""},
        {"key": "date",     "label": "日付",             "type": "text",     "default": ""},
        {"key": "author",   "label": "作成者・組織名",   "type": "text",     "default": ""},
    ],
    "section": [
        {"key": "section_number", "label": "章番号",               "type": "text",     "default": "01"},
        {"key": "section_title",  "label": "セクションタイトル",   "type": "text",     "default": ""},
        {"key": "description",    "label": "説明文（任意）",        "type": "textarea", "default": ""},
    ],
    "kpi3": [
        {"key": "slide_title",  "label": "スライドタイトル", "type": "text", "default": ""},
        {"key": "kpi1_label",   "label": "指標1 ラベル",     "type": "text", "default": "指標1"},
        {"key": "kpi1_value",   "label": "指標1 数値",       "type": "text", "default": "—"},
        {"key": "kpi1_unit",    "label": "指標1 単位",       "type": "text", "default": ""},
        {"key": "kpi1_note",    "label": "指標1 補足",       "type": "text", "default": ""},
        {"key": "kpi2_label",   "label": "指標2 ラベル",     "type": "text", "default": "指標2"},
        {"key": "kpi2_value",   "label": "指標2 数値",       "type": "text", "default": "—"},
        {"key": "kpi2_unit",    "label": "指標2 単位",       "type": "text", "default": ""},
        {"key": "kpi2_note",    "label": "指標2 補足",       "type": "text", "default": ""},
        {"key": "kpi3_label",   "label": "指標3 ラベル",     "type": "text", "default": "指標3"},
        {"key": "kpi3_value",   "label": "指標3 数値",       "type": "text", "default": "—"},
        {"key": "kpi3_unit",    "label": "指標3 単位",       "type": "text", "default": ""},
        {"key": "kpi3_note",    "label": "指標3 補足",       "type": "text", "default": ""},
    ],
    "bullet": [
        {"key": "slide_title", "label": "スライドタイトル",       "type": "text",     "default": ""},
        {"key": "lead",        "label": "リード文（任意）",        "type": "text",     "default": ""},
        {"key": "bullets",     "label": "箇条書き（1行1項目）",   "type": "textarea", "default": ""},
    ],
    "two_column": [
        {"key": "slide_title",   "label": "スライドタイトル",   "type": "text",     "default": ""},
        {"key": "left_title",    "label": "左カラムタイトル",   "type": "text",     "default": ""},
        {"key": "left_content",  "label": "左カラム内容",       "type": "textarea", "default": ""},
        {"key": "right_title",   "label": "右カラムタイトル",   "type": "text",     "default": ""},
        {"key": "right_content", "label": "右カラム内容",       "type": "textarea", "default": ""},
    ],
    "chart_bar": [
        {"key": "slide_title",  "label": "スライドタイトル",                 "type": "text",     "default": ""},
        {"key": "chart_title",  "label": "グラフタイトル",                   "type": "text",     "default": ""},
        {"key": "data_csv",     "label": "データ（ラベル,値 — 1行1データ）", "type": "textarea", "default": "2021,100\n2022,120\n2023,145\n2024,165"},
        {"key": "y_label",      "label": "Y軸ラベル（任意）",                "type": "text",     "default": ""},
        {"key": "note",         "label": "注記（任意）",                     "type": "text",     "default": ""},
    ],
    "chart_line": [
        {"key": "slide_title",  "label": "スライドタイトル",                 "type": "text",     "default": ""},
        {"key": "chart_title",  "label": "グラフタイトル",                   "type": "text",     "default": ""},
        {"key": "data_csv",     "label": "データ（ラベル,値 — 1行1データ）", "type": "textarea", "default": "2021,100\n2022,120\n2023,145\n2024,165"},
        {"key": "y_label",      "label": "Y軸ラベル（任意）",                "type": "text",     "default": ""},
        {"key": "note",         "label": "注記（任意）",                     "type": "text",     "default": ""},
    ],
    "chart_pie": [
        {"key": "slide_title",  "label": "スライドタイトル",                 "type": "text",     "default": ""},
        {"key": "chart_title",  "label": "グラフタイトル",                   "type": "text",     "default": ""},
        {"key": "data_csv",     "label": "データ（ラベル,値 — 1行1データ）", "type": "textarea", "default": "カテゴリA,35\nカテゴリB,28\nカテゴリC,20\nその他,17"},
        {"key": "note",         "label": "注記（任意）",                     "type": "text",     "default": ""},
    ],
    "table": [
        {"key": "slide_title", "label": "スライドタイトル",              "type": "text",     "default": ""},
        {"key": "data_csv",    "label": "データ（1行目がヘッダー行）",   "type": "textarea", "default": "項目,2022年,2023年,2024年\n売上（百万円）,100,120,145\n利益（百万円）,10,15,22\n従業員数,50,55,62"},
        {"key": "note",        "label": "注記（任意）",                  "type": "text",     "default": ""},
    ],
    "quote": [
        {"key": "message", "label": "メインメッセージ（インパクトのある一文）", "type": "textarea", "default": ""},
        {"key": "source",  "label": "出典・補足（任意）",                       "type": "text",     "default": ""},
    ],
    "closing": [
        {"key": "main_message", "label": "メインメッセージ",                     "type": "text",     "default": "ありがとうございました"},
        {"key": "contact",      "label": "連絡先・URL（任意）",                  "type": "textarea", "default": ""},
        {"key": "next_steps",   "label": "次のステップ（任意、1行1項目）",       "type": "textarea", "default": ""},
    ],
}

BUILDERS = {
    "title":      s_title,
    "section":    s_section,
    "kpi3":       s_kpi3,
    "bullet":     s_bullet,
    "two_column": s_two_column,
    "chart_bar":  s_chart_bar,
    "chart_line": s_chart_line,
    "chart_pie":  s_chart_pie,
    "table":      s_table,
    "quote":      s_quote,
    "closing":    s_closing,
}

# ── Main entry ─────────────────────────────────────────────────

def build_pptx(
    slides: list,
    theme: Theme,
    company: str = "",
    logo_bytes: bytes | None = None,
    logo_in_header: bool = True,
) -> io.BytesIO:
    """
    Build a PPTX file from a list of slide dicts and a Theme.
    logo_bytes      : raw bytes of a PNG/JPG logo image (optional)
    logo_in_header  : if True, logo appears in every slide's title bar;
                      it always appears prominently on the title slide.
    """
    prs = Presentation()
    prs.slide_width  = Inches(W)
    prs.slide_height = Inches(H)
    for pg, slide_data in enumerate(slides, start=1):
        stype = slide_data.get("type", "bullet")
        builder = BUILDERS.get(stype, s_bullet)
        # Title slide always gets the logo; other slides only if logo_in_header
        show_logo = logo_bytes if (stype == "title" or logo_in_header) else None
        builder(prs, slide_data, theme, pg, company, show_logo)
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf
